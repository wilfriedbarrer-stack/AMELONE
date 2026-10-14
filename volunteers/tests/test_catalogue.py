import csv
import io
from datetime import date, timedelta
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.core import signing
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import Workbook

from volunteers.catalogue_import import (
    MAX_FILE_BYTES,
    apply_catalogue_preview,
    parse_catalogue_upload,
    prepare_catalogue_preview,
)
from volunteers.models import Beverage, Event, EventDay, MealOrder, Member, Menu, Registration
from volunteers.services import confirm_assignment, create_assignment
from volunteers.views_catalogue import PREVIEW_SALT, PREVIEW_SESSION_KEY

from .factories import DataMixin


def csv_upload(content, name="catalogue.csv"):
    return SimpleUploadedFile(name, content.encode("utf-8-sig"), content_type="text/csv")


def xlsx_upload(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile("catalogue.xlsx", output.getvalue())


class CatalogueParserTests(SimpleTestCase):
    def test_excel_compatible_csv_accepts_bom_french_decimals_and_meal_types(self):
        preview = parse_catalogue_upload(csv_upload(
            'nom;description;supplément_euros;type\nRepas végétarien;Plat et dessert;2,50;repas\nGoûter;Fruit;0;collation\n'
        ), "menus")
        self.assertTrue(preview.valid)
        self.assertEqual([row["surcharge_cents"] for row in preview.rows], [250, 0])
        self.assertEqual([row["meal_type"] for row in preview.rows], ["full_meal", "snack"])

    def test_csv_comma_separator_and_quoted_description(self):
        preview = parse_catalogue_upload(csv_upload('nom,description,supplement_euros\nEau,"Eau, plate",1.50\n'), "boissons")
        self.assertTrue(preview.valid)
        self.assertEqual(preview.rows[0]["description"], "Eau, plate")

    def test_extra_values_in_comma_csv_show_the_bad_row_without_hiding_other_rows(self):
        preview = parse_catalogue_upload(csv_upload("nom,supplement_euros\nEau,1\nJus,2,extra\n"), "boissons")
        self.assertEqual(len(preview.rows), 2)
        self.assertEqual(preview.rows[0]["errors"], [])
        self.assertFalse(preview.valid)
        self.assertIn("plus de valeurs", " ".join(preview.rows[1]["errors"]))

    def test_all_bad_rows_are_shown_and_duplicate_names_block_import(self):
        preview = parse_catalogue_upload(csv_upload(
            'nom;supplement_euros;type\nMenu;-2;autre\nmenu;3.001;repas\n;1;collation\n'
        ), "menus")
        self.assertFalse(preview.valid)
        self.assertEqual(len(preview.rows), 3)
        self.assertTrue(all(row["errors"] for row in preview.rows))
        self.assertIn("plusieurs fois", " ".join(preview.rows[1]["errors"]))

    def test_single_name_column_is_valid_for_beverages(self):
        preview = parse_catalogue_upload(csv_upload("nom\nEau\nJus de pomme\n"), "boissons")
        self.assertTrue(preview.valid)
        self.assertEqual(len(preview.rows), 2)

    def test_invalid_encoding_unknown_columns_and_empty_upload_are_rejected(self):
        uploads = [
            SimpleUploadedFile("catalogue.csv", "nom\nThé\n".encode("cp1252")),
            csv_upload("nom;url\nEau;https://example.test\n"),
            csv_upload(""),
        ]
        for upload in uploads:
            with self.subTest(upload=upload):
                self.assertFalse(parse_catalogue_upload(upload, "boissons").valid)

    def test_size_and_row_limits_are_enforced(self):
        large = SimpleUploadedFile("catalogue.csv", b"x" * (MAX_FILE_BYTES + 1))
        self.assertIn("2 Mo", parse_catalogue_upload(large, "menus").errors[0])
        rows = "nom\n" + "\n".join(f"Boisson {number}" for number in range(501))
        preview = parse_catalogue_upload(csv_upload(rows), "boissons")
        self.assertFalse(preview.valid)
        self.assertEqual(len(preview.rows), 500)

    def test_xlsx_is_supported_and_formulas_are_rejected(self):
        upload = xlsx_upload([["nom", "supplement_euros", "type"], ["Plat", 1.5, "repas"], ["Goûter", 0, "collation"]])
        preview = parse_catalogue_upload(upload, "menus")
        self.assertTrue(preview.valid)
        self.assertEqual(preview.rows[0]["surcharge_cents"], 150)
        formula = xlsx_upload([["nom", "supplement_euros"], ["Eau", "=1+1"]])
        preview = parse_catalogue_upload(formula, "boissons")
        self.assertFalse(preview.valid)
        self.assertIn("formules Excel", " ".join(preview.rows[0]["errors"]))

    def test_xlsx_decompression_bomb_and_xml_entities_are_rejected(self):
        for contents in (b"X" * 500000, b'<!DOCTYPE x [<!ENTITY injected "text">]><x/>'):
            output = io.BytesIO()
            with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
                archive.writestr("xl/workbook.xml", contents)
            preview = parse_catalogue_upload(SimpleUploadedFile("catalogue.xlsx", output.getvalue()), "menus")
            self.assertFalse(preview.valid)
            self.assertTrue(preview.errors)

    def test_xlsx_extra_cells_beyond_allowed_columns_are_not_silently_ignored(self):
        workbook = Workbook()
        workbook.active.append(["nom"])
        workbook.active.cell(row=2, column=1, value="Eau")
        workbook.active.cell(row=2, column=40, value="Donnée non prise en charge")
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        preview = parse_catalogue_upload(SimpleUploadedFile("catalogue.xlsx", output.getvalue()), "boissons")
        self.assertFalse(preview.valid)
        self.assertIn("20 colonnes", " ".join(preview.errors))


class CatalogueDatabaseTests(DataMixin, TestCase):
    def preview(self, content, kind="menus"):
        return prepare_catalogue_preview(parse_catalogue_upload(csv_upload(content), kind), self.day, kind)

    def test_import_is_atomic_and_existing_updates_require_confirmation(self):
        existing = Menu.objects.create(event_day=self.day, name="Menu existant", surcharge_cents=100)
        preview = self.preview("nom;supplement_euros;type\nNouveau;1;repas\nMenu existant;2;repas\n")
        self.assertTrue(preview.valid)
        with self.assertRaisesMessage(ValidationError, "Confirmez explicitement"):
            apply_catalogue_preview(day=self.day, kind="menus", rows=preview.rows)
        self.assertEqual(Menu.objects.count(), 1)
        existing.refresh_from_db()
        self.assertEqual(existing.surcharge_cents, 100)
        self.assertEqual(apply_catalogue_preview(day=self.day, kind="menus", rows=preview.rows, confirm_updates=True), 2)
        existing.refresh_from_db()
        self.assertEqual(existing.surcharge_cents, 200)

    def test_edits_since_preview_are_detected_without_partial_import(self):
        existing = Beverage.objects.create(event_day=self.day, name="Eau")
        preview = self.preview("nom;supplement_euros\nJus;1\nEau;2\n", "boissons")
        existing.description = "Modification concurrente"
        existing.save()
        with self.assertRaisesMessage(ValidationError, "changé depuis l'aperçu"):
            apply_catalogue_preview(day=self.day, kind="boissons", rows=preview.rows, confirm_updates=True)
        self.assertEqual(Beverage.objects.count(), 1)

    def test_menu_type_update_cannot_invalidate_a_reserved_full_day_meal(self):
        menu = Menu.objects.create(event_day=self.day, name="Repas")
        MealOrder.objects.create(registration=self.registration(), wants_meal=True, menu=menu)
        preview = self.preview("nom;type\nRepas;collation\n")
        self.assertFalse(preview.valid)
        self.assertIn("incompatible", " ".join(preview.rows[0]["errors"]))

    def test_preserved_legacy_order_does_not_block_price_or_description_update(self):
        registration = self.registration()
        menu = Menu.objects.create(event_day=self.day, name="Repas")
        order = MealOrder.objects.create(registration=registration, wants_meal=True, menu=menu)
        # Mimic a historical half-day meal preserved by the data migration.
        Registration.objects.filter(pk=registration.pk).update(availability="morning")
        MealOrder.objects.filter(pk=order.pk).update(legacy_availability="morning")
        preview = self.preview("nom;description;supplement_euros\nRepas;Description améliorée;2\n")
        self.assertTrue(preview.valid)
        apply_catalogue_preview(day=self.day, kind="menus", rows=preview.rows, confirm_updates=True)
        menu.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(menu.surcharge_cents, 200)
        self.assertEqual(menu.meal_type, "full_meal")
        self.assertEqual(order.total_surcharge_cents, 0)


class CatalogueViewsTests(DataMixin, TestCase):
    def setUp(self):
        self.client.force_login(self.organizer)

    def upload(self, content, kind="menus", day=None):
        return self.client.post(reverse("catalogue_import", args=[self.event.pk, kind]), {
            "action": "preview", "event_day": (day or self.day).pk, "file": csv_upload(content),
        })

    def confirm(self, response, **extra):
        return self.client.post(reverse("catalogue_import", args=[self.event.pk, "menus"]), {
            "action": "confirm", "preview_token": response.context["confirm_form"].initial["preview_token"], **extra,
        })

    def test_upload_only_previews_and_confirmation_uses_server_data(self):
        response = self.upload("nom;supplement_euros;type\nRepas végétarien;2,50;repas\nCollation;0;collation\n")
        self.assertContains(response, "Aperçu des 2 lignes")
        self.assertEqual(Menu.objects.count(), 0)
        response = self.confirm(response, rows='[{"name":"Pirate","surcharge_cents":0}]')
        self.assertRedirects(response, reverse("catering", args=[self.event.pk]))
        self.assertEqual(Menu.objects.count(), 2)
        self.assertFalse(Menu.objects.filter(name="Pirate").exists())
        self.assertNotIn(PREVIEW_SESSION_KEY, self.client.session)

    def test_invalid_rows_prevent_confirmation_of_entire_file(self):
        response = self.upload("nom;supplement_euros\nBon;2\nMauvais;-1\n")
        self.assertIsNone(response.context["confirm_form"])
        self.assertContains(response, "Aucun changement n'a été enregistré")
        self.assertEqual(Menu.objects.count(), 0)

    def test_duplicate_name_update_requires_checkbox_and_preserves_other_options(self):
        menu = Menu.objects.create(event_day=self.day, name="Repas", surcharge_cents=50, is_active=False)
        other = Menu.objects.create(event_day=self.day, name="Conservé")
        response = self.upload("nom;supplement_euros\nRepas;2\n")
        response = self.confirm(response)
        self.assertContains(response, "Confirmez explicitement")
        menu.refresh_from_db()
        self.assertEqual(menu.surcharge_cents, 50)
        token = response.context["confirm_form"]["preview_token"].value()
        response = self.client.post(reverse("catalogue_import", args=[self.event.pk, "menus"]), {
            "action": "confirm", "preview_token": token, "confirm_updates": "on",
        })
        self.assertEqual(response.status_code, 302)
        menu.refresh_from_db()
        self.assertEqual(menu.surcharge_cents, 200)
        self.assertFalse(menu.is_active)
        self.assertTrue(Menu.objects.filter(pk=other.pk).exists())

    def test_day_from_other_event_cannot_receive_import(self):
        other_event = Event.objects.create(
            name="Autre", location="Club", registration_opens_at=self.event.registration_opens_at,
            registration_closes_at=self.event.registration_closes_at, meal_orders_close_at=self.event.meal_orders_close_at,
        )
        other_day = EventDay.objects.create(event=other_event, date=self.day.date)
        response = self.upload("nom\nRepas\n", day=other_day)
        self.assertContains(response, "Sélectionnez un choix valide")
        self.assertEqual(Menu.objects.count(), 0)

    def test_tampered_preview_token_and_other_session_are_rejected(self):
        response = self.upload("nom\nRepas\n")
        valid_token = response.context["confirm_form"].initial["preview_token"]
        for token in (valid_token + "broken", signing.dumps({"nonce": "different"}, salt=PREVIEW_SALT)):
            response = self.client.post(reverse("catalogue_import", args=[self.event.pk, "menus"]), {
                "action": "confirm", "preview_token": token,
            })
            self.assertEqual(response.status_code, 200)
            self.assertEqual(Menu.objects.count(), 0)

    def test_preview_expires_and_successful_confirmation_cannot_be_replayed(self):
        response = self.upload("nom\nRepas\n")
        token = response.context["confirm_form"].initial["preview_token"]
        with patch("django.core.signing.time.time", return_value=9999999999):
            expired = self.confirm(response)
        self.assertContains(expired, "aperçu a expiré")
        self.assertEqual(Menu.objects.count(), 0)
        self.assertEqual(self.confirm(response).status_code, 302)
        replay = self.client.post(reverse("catalogue_import", args=[self.event.pk, "menus"]), {
            "action": "confirm", "preview_token": token,
        })
        self.assertContains(replay, "plus disponible")
        self.assertEqual(Menu.objects.count(), 1)

    def test_catalogue_and_templates_are_organizer_only(self):
        for user in (None, self.family_user):
            self.client.logout()
            if user:
                self.client.force_login(user)
            for url in (
                reverse("catalogue_import", args=[self.event.pk, "menus"]),
                reverse("catalogue_template", args=["boissons"]),
            ):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 403 if user else 302)


class RecapReportsTests(DataMixin, TestCase):
    def setUp(self):
        self.client.force_login(self.organizer)

    def test_post_report_includes_assignments_and_unassigned_participants_only_for_event(self):
        registration = self.registration()
        assignment = create_assignment(registration=registration, station_slot=self.morning_slot, actor=self.organizer)
        confirm_assignment(assignment=assignment, confirmed_by=self.organizer)
        self.registration(member=self.other_member, availability=Registration.Availability.MORNING)
        response = self.client.get(reverse("recap_assignments", args=[self.event.pk]))
        self.assertContains(response, "Accueil")
        self.assertContains(response, "Confirmée")
        self.assertContains(response, self.other_member.full_name)
        self.assertContains(response, "Non affecté")
        self.assertEqual(response.context["unassigned_count"], 1)
        self.assertEqual(response.context["participant_count"], 2)

    def test_meal_report_distinguishes_ordered_none_and_missing_choices(self):
        first = self.registration()
        second = self.registration(member=self.other_member, availability=Registration.Availability.MORNING)
        third_member = Member.objects.create(family=self.family, first_name="Troisième", last_name="Durand", birth_date=date(1990, 1, 1))
        third = self.registration(member=third_member, availability=Registration.Availability.AFTERNOON)
        cancelled_member = Member.objects.create(family=self.family, first_name="Annulé", last_name="Durand", birth_date=date(1990, 1, 1))
        self.registration(member=cancelled_member, status=Registration.Status.CANCELLED)
        menu = Menu.objects.create(event_day=self.day, name="Plat végétarien", surcharge_cents=125)
        beverage = Beverage.objects.create(event_day=self.day, name="Jus de pomme")
        MealOrder.objects.create(registration=first, wants_meal=True, menu=menu, beverage=beverage, payment_method="sumup")
        MealOrder.objects.create(registration=second, wants_meal=False)
        response = self.client.get(reverse("recap_meals", args=[self.event.pk]))
        self.assertContains(response, "Plat végétarien")
        self.assertContains(response, "Paiement via lien SumUp")
        self.assertContains(response, "Pas de repas")
        self.assertContains(response, "Non renseigné")
        self.assertContains(response, third.member.full_name)
        self.assertNotContains(response, cancelled_member.full_name)
        self.assertEqual(response.context["ordered_count"], 1)
        self.assertEqual(response.context["no_meal_count"], 1)
        self.assertEqual(response.context["missing_count"], 1)
        self.assertEqual(response.context["participant_count"], 3)
        self.assertEqual(response.context["quantities"][0]["count"], 1)

    def test_reports_and_exports_never_expose_private_family_data_to_public_or_family(self):
        self.registration()
        for user in (None, self.family_user):
            self.client.logout()
            if user:
                self.client.force_login(user)
            for name in ("recap_assignments", "recap_assignments_csv", "recap_meals", "recap_meals_csv"):
                response = self.client.get(reverse(name, args=[self.event.pk]))
                self.assertEqual(response.status_code, 403 if user else 302)
                self.assertNotIn(self.member.full_name.encode(), response.content)

    def test_csv_exports_escape_spreadsheet_formulas_and_include_no_meal_participants(self):
        self.family.name = "=1+1"
        self.family.save()
        self.member.first_name = "+Lina"
        self.member.save()
        self.registration()
        for name in ("recap_assignments_csv", "recap_meals_csv"):
            response = self.client.get(reverse(name, args=[self.event.pk]))
            rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
            self.assertEqual(rows[1][1], "'=1+1")
            self.assertEqual(rows[1][2], "'+Lina Durand")
            self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assertIn("Non renseigné", rows[1])

    def test_report_does_not_include_participants_from_another_event(self):
        self.registration()
        other = Event.objects.create(
            name="Autre événement", location="Club", registration_opens_at=self.event.registration_opens_at,
            registration_closes_at=self.event.registration_closes_at, meal_orders_close_at=self.event.meal_orders_close_at,
        )
        day = EventDay.objects.create(event=other, date=self.day.date + timedelta(days=1))
        Registration.objects.create(member=self.other_member, event_day=day, availability="morning")
        for name in ("recap_assignments", "recap_meals"):
            response = self.client.get(reverse(name, args=[self.event.pk]))
            self.assertContains(response, self.member.full_name)
            self.assertNotContains(response, self.other_member.full_name)
