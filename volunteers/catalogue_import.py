"""Validate catalogue uploads before writing any menu or beverage to the database."""

import csv
import hashlib
import io
import json
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from django.db import transaction

from .models import Beverage, EventDay, Menu, Registration

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
MAX_DESCRIPTION_LENGTH = 4000
MAX_XLSX_UNCOMPRESSED_BYTES = 12 * 1024 * 1024
CATALOGUE_KINDS = {"menus": Menu, "boissons": Beverage}
MEAL_TYPES = {
    "full_meal": "full_meal",
    "repas": "full_meal",
    "repas complet": "full_meal",
    "snack": "snack",
    "collation": "snack",
}


@dataclass
class CataloguePreview:
    rows: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def valid(self):
        return bool(self.rows) and not self.errors and all(not row["errors"] for row in self.rows)


def _normalize_header(value):
    decomposed = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(character for character in decomposed if not unicodedata.combining(character))


def _csv_rows(raw):
    try:
        decoded = raw.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValidationError("Le fichier CSV doit être enregistré en UTF-8.") from error
    if "\x00" in decoded:
        raise ValidationError("Le fichier CSV contient des caractères non valides.")
    try:
        # Detect from the header, so malformed data rows remain visible in preview.
        header_line = decoded.split("\n", 1)[0]
        dialect = csv.Sniffer().sniff(header_line, delimiters=";,\t")
    except csv.Error:
        # A single-column beverage file (nom) is also a valid catalogue.
        dialect = None
    reader = (
        csv.reader(io.StringIO(decoded, newline=""), dialect, strict=True)
        if dialect else csv.reader(io.StringIO(decoded, newline=""), delimiter=";", strict=True)
    )
    try:
        for row in reader:
            yield reader.line_num, row, False
    except csv.Error as error:
        raise ValidationError("Le fichier CSV n'est pas correctement formé.") from error


def _xlsx_rows(raw):
    # Limit both compressed and expanded size before openpyxl touches the archive.
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 300 or sum(item.file_size for item in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValidationError("Le classeur Excel est trop volumineux après décompression.")
            if len({item.filename for item in members}) != len(members):
                raise ValidationError("Le classeur Excel contient des fichiers en double.")
            for item in members:
                if item.flag_bits & 1 or item.file_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise ValidationError("Les classeurs Excel chiffrés ne sont pas acceptés.")
                if item.file_size > max(item.compress_size, 1) * 200:
                    raise ValidationError("Le taux de compression du classeur Excel est trop élevé.")
                if item.filename.lower().endswith(".xml"):
                    data = archive.read(item)
                    if b"\x00" in data or re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", data, re.IGNORECASE):
                        raise ValidationError("Le classeur Excel contient des données XML non autorisées.")
                    if item.filename.startswith("xl/worksheets/"):
                        try:
                            for _, element in ElementTree.iterparse(io.BytesIO(data), events=("end",)):
                                if element.tag.rsplit("}", 1)[-1] == "c":
                                    reference = element.attrib.get("r", "")
                                    match = re.fullmatch(r"([A-Z]+)(\d+)", reference)
                                    if not match:
                                        raise ValidationError("Le classeur contient une référence de cellule invalide.")
                                    column = 0
                                    for character in match[1]:
                                        column = column * 26 + ord(character) - ord("A") + 1
                                    if column > 20 or int(match[2]) > 1001:
                                        raise ValidationError("Le classeur est limité à 20 colonnes et 1 000 lignes, dont 500 lignes de données.")
                                element.clear()
                        except ElementTree.ParseError as error:
                            raise ValidationError("Le fichier Excel est illisible ou endommagé.") from error
            if "xl/workbook.xml" not in archive.namelist():
                raise ValidationError("Ce fichier n'est pas un classeur Excel .xlsx.")
    except (BadZipFile, OSError, RuntimeError) as error:
        raise ValidationError("Le fichier Excel est illisible ou endommagé.") from error

    from openpyxl import load_workbook

    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        if len(workbook.worksheets) != 1:
            raise ValidationError("Utilisez un classeur avec une seule feuille de menus ou de boissons.")
        sheet = workbook.worksheets[0]
        # Ignore advertised worksheet dimensions; they can hide rows or be forged.
        sheet.reset_dimensions()
        for number, cells in enumerate(sheet.iter_rows(max_col=21), start=1):
            values = [cell.value for cell in cells]
            while values and values[-1] is None:
                values.pop()
            yield number, values, any(cell.data_type == "f" for cell in cells)
    except ValidationError:
        raise
    except Exception as error:
        raise ValidationError("Le fichier Excel est illisible ou endommagé.") from error
    finally:
        if workbook is not None:
            workbook.close()


def _parse_amount(value):
    raw = str(value if value is not None else "").strip().replace(",", ".")
    if not raw:
        return 0
    if not re.fullmatch(r"\d+(?:\.\d{1,2})?", raw):
        raise ValidationError("Le supplément doit être un montant positif avec au plus deux décimales.")
    try:
        amount = Decimal(raw)
        if amount > Decimal("10000"):
            raise ValidationError("Le supplément doit être inférieur ou égal à 10 000 €.")
        return int(amount * 100)
    except InvalidOperation as error:
        raise ValidationError("Le supplément est invalide.") from error


def parse_catalogue_upload(upload, kind, *, identity_by_type=False):
    if kind not in CATALOGUE_KINDS:
        raise ValidationError("Choisissez un import de menus ou de boissons.")
    preview = CataloguePreview()
    if upload.size > MAX_FILE_BYTES:
        preview.errors.append("Le fichier ne doit pas dépasser 2 Mo.")
        return preview
    raw = upload.read(MAX_FILE_BYTES + 1)
    if len(raw) > MAX_FILE_BYTES:
        preview.errors.append("Le fichier ne doit pas dépasser 2 Mo.")
        return preview
    suffix = Path(upload.name).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        preview.errors.append("Choisissez un fichier CSV UTF-8 ou Excel .xlsx.")
        return preview
    try:
        rows = iter(_csv_rows(raw) if suffix == ".csv" else _xlsx_rows(raw))
        _, headers, header_formula = next(rows)
        headers = [_normalize_header(value) for value in headers]
        allowed = {"nom", "description", "supplement_euros"}
        if kind == "menus":
            allowed.add("type")
        if header_formula or len(headers) > 20 or "nom" not in headers:
            raise ValidationError("La première ligne doit contenir les titres de colonnes, dont « nom ».")
        if len(set(headers)) != len(headers) or set(headers) - allowed:
            raise ValidationError("Colonnes autorisées : " + ", ".join(sorted(allowed)) + ".")
        seen = set()
        for scanned_count, (number, values, has_formula) in enumerate(rows, start=1):
            if scanned_count > 1000:
                raise ValidationError("Le fichier est limité à 1 000 lignes, dont 500 lignes de données.")
            if not any(value is not None and str(value).strip() for value in values):
                continue
            if len(preview.rows) >= MAX_ROWS:
                raise ValidationError("Le fichier ne doit pas contenir plus de 500 lignes de données.")
            row = {
                "line": number,
                "name": "",
                "description": "",
                "surcharge_cents": 0,
                "meal_type": "full_meal",
                "errors": [],
            }
            padded = values + [""] * max(len(headers) - len(values), 0)
            data = dict(zip(headers, padded))
            if has_formula:
                row["errors"].append("Les formules Excel ne sont pas autorisées.")
            if len(values) > len(headers):
                row["errors"].append("Cette ligne contient plus de valeurs que de colonnes.")
            row["name"] = str(data.get("nom") or "").strip()
            row["description"] = str(data.get("description") or "").strip()
            if not row["name"] or len(row["name"]) > 160:
                row["errors"].append("Le nom est obligatoire et limité à 160 caractères.")
            if len(row["description"]) > MAX_DESCRIPTION_LENGTH:
                row["errors"].append("La description est limitée à 4 000 caractères.")
            try:
                row["surcharge_cents"] = _parse_amount(data.get("supplement_euros"))
            except ValidationError as error:
                row["errors"].extend(error.messages)
            if kind == "menus":
                meal_type = str(data.get("type") or "full_meal").strip().lower()
                if meal_type not in MEAL_TYPES:
                    row["errors"].append("Le type doit être « repas » ou « collation ».")
                else:
                    row["meal_type"] = MEAL_TYPES[meal_type]
            identity = (
                (row["name"].casefold(), row["meal_type"])
                if identity_by_type and kind == "menus" else row["name"].casefold()
            )
            if identity in seen:
                row["errors"].append(
                    "Ce nom et ce type apparaissent plusieurs fois dans le fichier."
                    if identity_by_type and kind == "menus"
                    else "Ce nom apparaît plusieurs fois dans le fichier."
                )
            seen.add(identity)
            preview.rows.append(row)
    except StopIteration:
        preview.errors.append("Le fichier est vide.")
    except ValidationError as error:
        preview.errors.extend(error.messages)
    if not preview.rows and not preview.errors:
        preview.errors.append("Le fichier ne contient aucun menu ou boisson.")
    return preview


def _fingerprint(option):
    if option is None:
        return None
    state = {
        "id": str(option.pk),
        "name": option.name,
        "description": option.description,
        "surcharge_cents": option.surcharge_cents,
        "is_active": option.is_active,
        "updated_at": option.updated_at.isoformat(),
    }
    if isinstance(option, Menu):
        state["meal_type"] = option.meal_type
    return hashlib.sha256(json.dumps(state, sort_keys=True).encode()).hexdigest()


def _validate_option(option):
    errors = []
    try:
        option.full_clean()
    except ValidationError as error:
        errors.extend(error.messages)
    if isinstance(option, Menu) and not option._state.adding:
        previous_type = Menu.objects.filter(pk=option.pk).values_list("meal_type", flat=True).first()
        if option.meal_type == previous_type:
            return errors
        orders = option.orders.filter(wants_meal=True, registration__status=Registration.Status.ACTIVE)
        incompatible = (
            orders.exclude(registration__availability=Registration.Availability.FULL_DAY)
            if option.meal_type == "full_meal"
            else orders.filter(registration__availability=Registration.Availability.FULL_DAY)
        )
        if incompatible.exists():
            errors.append("Ce type de menu est incompatible avec des réservations existantes.")
    return errors


def _options_by_name(day, kind, lock=False):
    query = CATALOGUE_KINDS[kind].objects.filter(event_day=day)
    if lock:
        query = query.select_for_update(of=("self",)).order_by("pk")
    index = {}
    for option in query:
        identity = (
            (option.name.casefold(), option.meal_type)
            if day is None and kind == "menus" else option.name.casefold()
        )
        index.setdefault(identity, []).append(option)
    return index


def _row_identity(row, day, kind):
    if day is None and kind == "menus":
        return row["name"].casefold(), row["meal_type"]
    return row["name"].casefold()


def _option_from_row(row, day, kind, existing=None):
    option = existing or CATALOGUE_KINDS[kind](event_day=day)
    option.name = row["name"]
    option.description = row["description"]
    option.surcharge_cents = row["surcharge_cents"]
    if kind == "menus":
        option.meal_type = row["meal_type"]
    return option


def prepare_catalogue_preview(preview, day, kind):
    index = _options_by_name(day, kind)
    for row in preview.rows:
        matching = index.get(_row_identity(row, day, kind), [])
        row["action"] = "update" if matching else "create"
        row["expected"] = _fingerprint(matching[0]) if len(matching) == 1 else None
        if len(matching) > 1:
            row["errors"].append("Plusieurs options existantes portent ce nom : renommez-les avant l'import.")
        if not row["errors"]:
            option = _option_from_row(row, day, kind, matching[0] if matching else None)
            row["errors"].extend(_validate_option(option))
    return preview


@transaction.atomic
def apply_catalogue_preview(*, day, kind, rows, confirm_updates=False, actor=None):
    if day is None or actor is not None:
        # Local import callers remain compatible; shared mutations always reload
        # and lock the active organizer before touching catalogue rows.
        from .shared_catalogue import authorize_catalogue_actor

        authorize_catalogue_actor(actor)
    # Lock the day for local imports and all existing options in either scope.
    if day is not None:
        day = EventDay.objects.select_for_update().get(pk=day.pk)
    index = _options_by_name(day, kind, lock=True)
    options = []
    if not rows or len(rows) > MAX_ROWS:
        raise ValidationError("L'aperçu est vide ou invalide. Recommencez l'import.")
    for row in rows:
        if row.get("errors"):
            raise ValidationError("Corrigez toutes les erreurs avant de confirmer l'import.")
        matching = index.get(_row_identity(row, day, kind), [])
        if len(matching) > 1 or _fingerprint(matching[0] if matching else None) != row["expected"]:
            raise ValidationError("Le catalogue a changé depuis l'aperçu. Recommencez l'import.")
        existing = matching[0] if matching else None
        if existing and not confirm_updates:
            raise ValidationError("Confirmez explicitement la mise à jour des options existantes.")
        option = _option_from_row(row, day, kind, existing)
        errors = _validate_option(option)
        if errors:
            raise ValidationError(errors)
        options.append(option)
    # An empty shared catalogue offers no rows to lock. Stable insertion order
    # also avoids crossed unique-index waits between two concurrent imports.
    options.sort(key=lambda option: (option.name.casefold(), getattr(option, "meal_type", "")))
    for option in options:
        option.save()
    return len(options)
